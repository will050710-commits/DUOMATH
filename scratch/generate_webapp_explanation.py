from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

out_dir = Path(r"D:\Downloads")
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / "Duosteam_webapp_full_feature_explanation.docx"

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

heading = doc.add_heading('Giải thích toàn bộ tính năng của web app DuoSteam', level=1)
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('Tài liệu này giải thích toàn bộ hệ thống web app, không chỉ một bài học riêng lẻ. Nó tập trung vào các tính năng chính, luồng hoạt động và cách các module kết nối với nhau.')

doc.add_paragraph('')

# 1 Overview
p = doc.add_paragraph(); p.add_run('1. Tổng quan về web app').bold = True
p = doc.add_paragraph('DuoSteam là một nền tảng học tập trực tuyến dành cho học sinh, tập trung vào môn Toán. Web app có mục tiêu cung cấp bài học tương tác, câu hỏi luyện tập, theo dõi kết quả, gamification và trải nghiệm người dùng mượt mà.')

# 2 Architecture
p = doc.add_paragraph(); p.add_run('2. Kiến trúc tổng thể').bold = True
for item in [
    'Frontend được xây dựng bằng Next.js, dùng React để render giao diện.',
    'Các component được chia theo chức năng: trang chủ, lesson, sidebar, auth, gamification, thống kê, phản hồi và kết quả.',
    'Context API được dùng để quản lý trạng thái như auth và dữ liệu học tập.',
    'Một số tính năng dùng hook riêng để xử lý logic như gamification hoặc thống kê.',
    'Ứng dụng có thể hiển thị nội dung theo nhiều lesson khác nhau trong cùng một hệ thống chung.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 3 Main features
p = doc.add_paragraph(); p.add_run('3. Các tính năng chính của web app').bold = True
for item in [
    'Trang chủ và điều hướng tổng quát.',
    'Hiển thị danh sách bài học theo từng chương và từng lớp.',
    'Bài học có nội dung lý thuyết, video, câu hỏi tương tác và phần thực hành.',
    'Chuyển đổi ngôn ngữ giữa tiếng Việt và tiếng Anh.',
    'Mini-game luyện tập với nhiều loại câu hỏi.',
    'Điều chỉnh độ khó câu hỏi theo mức nhận biết, thông hiểu, vận dụng và vận dụng cao.',
    'Gamification, phần thưởng và HUD thông báo tiến độ.',
    'Lưu kết quả làm bài của người dùng.',
    'Hỗ trợ công cụ toán học và dịch thuật.',
    'Thanh điều hướng, sidebar và chuyển trang mượt mà.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 4 Landing page
p = doc.add_paragraph(); p.add_run('4. Trang chủ').bold = True
for item in [
    'Trang chủ được render bởi component TrangChuForm.',
    'Nó đóng vai trò là cửa ngõ vào toàn bộ hệ thống, giúp người dùng chọn chức năng hoặc chuyển tới bài học.',
    'Trang này có thể dùng để giới thiệu giao diện học tập và điều hướng tới các phần nội dung chính.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 5 Layout system
p = doc.add_paragraph(); p.add_run('5. Hệ thống layout và điều hướng').bold = True
for item in [
    'LayoutClient kết hợp PageTransition để tạo hiệu ứng chuyển trang mượt mà.',
    'GlobalSidebar hiện ở mọi trang để điều hướng giữa các module chính.',
    'Cấu trúc này giúp web app có trải nghiệm thống nhất trên nhiều trang.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 6 Lesson engine
p = doc.add_paragraph(); p.add_run('6. Hệ thống lesson engine').bold = True
for item in [
    'PremiumLessonEngine là core của trải nghiệm lesson.',
    'Nó chịu trách nhiệm hiển thị bố cục chung của một bài học: tiêu đề, mục tiêu, nội dung lý thuyết, video, mini-game.',
    'Nó cũng quản lý trạng thái câu hỏi và điểm số.',
    'Các lesson cụ thể như Lesson2_TapHop chỉ cung cấp nội dung và dữ liệu; phần render và logic chung do engine xử lý.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 7 Interactive learning
p = doc.add_paragraph(); p.add_run('7. Tính năng học tập tương tác').bold = True
for item in [
    'Trắc nghiệm: người dùng chọn đáp án, xem giải thích và nhận điểm.',
    'Đúng/Sai: người dùng đánh giá mệnh đề và xem lời giải thích ngay sau đó.',
    'Điền từ: người dùng nhập đáp án và nhấn kiểm tra.',
    'Mỗi chế độ đều có kết quả cuối cùng và nút chơi lại.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 8 Difficulty and gamification
p = doc.add_paragraph(); p.add_run('8. Độ khó và gamification').bold = True
for item in [
    'Người dùng có thể chọn độ khó thủ công hoặc để hệ thống tự điều chỉnh.',
    'Hệ thống dùng mức độ NB, TH, VD, VDC để phân loại câu hỏi.',
    'Sau khi làm bài, dữ liệu sẽ được gửi tới hook gamification để tạo phản hồi và phần thưởng.',
    'HUD có thể hiện thông báo tiến độ và động lực học tập.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 9 Auth and persistence
p = doc.add_paragraph(); p.add_run('9. Xác thực người dùng và lưu dữ liệu').bold = True
for item in [
    'Web app dùng auth context để biết người dùng hiện tại.',
    'Khi kết thúc một bộ câu hỏi, hệ thống gọi saveGameResult để lưu tiến độ.',
    'Điều này cho phép theo dõi kết quả, điểm số và quá trình làm bài của từng người dùng.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 10 Extra tools
p = doc.add_paragraph(); p.add_run('10. Công cụ hỗ trợ học tập').bold = True
for item in [
    'DuoTranslate cung cấp chức năng dịch thuật trong giao diện.',
    'MathToolsPanel cung cấp các công cụ hữu ích cho toán học.',
    'LessonVideoPlayer cho phép nhúng video bài giảng và phụ đề.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 11 UX features
p = doc.add_paragraph(); p.add_run('11. Tính năng trải nghiệm người dùng').bold = True
for item in [
    'Có hiệu ứng reveal khi cuộn xuống.',
    'Nút bấm có animation và hiệu ứng phản hồi.',
    'Có thanh điều hướng sticky, bố cục 2 cột và responsive trên màn hình nhỏ.',
    'Giao diện có màu sắc hiện đại và chủ đề tối.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 12 Flow
p = doc.add_paragraph(); p.add_run('12. Luồng hoạt động của người dùng').bold = True
for item in [
    'Người dùng mở trang chủ và chọn bài học.',
    'Hệ thống điều hướng tới lesson tương ứng.',
    'Người dùng đọc lý thuyết, xem video và làm bài tập.',
    'Hệ thống cập nhật điểm số và lưu kết quả.',
    'Người dùng có thể tiếp tục học các lesson khác hoặc xem lại kết quả.',
]:
    para = doc.add_paragraph(style='List Bullet'); para.add_run(item)

# 13 Summary
p = doc.add_paragraph(); p.add_run('13. Kết luận').bold = True
p = doc.add_paragraph('Web app DuoSteam không chỉ là một trang học bài đơn lẻ, mà là một nền tảng học tập tương tác với nhiều module: trang chủ, lesson engine, mini-game, gamification, auth, thống kê và công cụ hỗ trợ. Đây là một hệ thống có cấu trúc rõ ràng, dễ mở rộng và phù hợp cho việc phát triển các trải nghiệm học tập trực tuyến chuyên nghiệp.')

# Save

doc.save(out_path)
print(f'Created: {out_path}')
print(f'Exists: {out_path.exists()}')
