from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

out_dir = Path(r"D:\Downloads")
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / "Lesson2_TapHop_code_explanation.docx"

# Create document

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# Title
doc.add_heading('Giải thích chi tiết code trang bài học Tập Hợp', level=1)
doc.add_paragraph('Tài liệu này được tạo dựa trên hai file chính: Lesson2_TapHop.js và PremiumLessonEngine.js.')
doc.add_paragraph('Ngày tạo: 2026-06-25')
doc.add_paragraph('')

# Overview
p = doc.add_paragraph()
p.add_run('1. Tổng quan về trang').bold = True
p = doc.add_paragraph()
p.add_run('Trang này là một bài học tương tác dành cho môn Toán 10, chủ đề “Tập hợp”. Nó kết hợp nhiều thành phần: phần giới thiệu, video bài giảng, lý thuyết, bài tập trắc nghiệm, đúng/sai, điền khuyết và một giao diện học tập có gamification.')

# Structure
p = doc.add_paragraph()
p.add_run('2. Cấu trúc của file chính').bold = True
p = doc.add_paragraph('File Lesson2_TapHop.js chịu trách nhiệm xây dựng nội dung bài học, bao gồm:')
for item in [
    'Các meta dữ liệu của bài học: tên bài, mục tiêu học tập, mục lục điều hướng.',
    'Danh sách câu hỏi trắc nghiệm, đúng/sai và điền khuyết.',
    'Nội dung lý thuyết chia thành các phần: Khởi động, Video, phần lý thuyết về tập hợp, tập con, hai tập bằng nhau, thực hành.',
    'Các component nhỏ dùng để tạo khung UI cho từng phần.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

# Main components
doc.add_paragraph('')
p = doc.add_paragraph()
p.add_run('3. Các component quan trọng trong Lesson2_TapHop.js').bold = True
for item in [
    'SectionHeader: tạo tiêu đề section với icon và đường kẻ phân cách.',
    'TheoryBlock: khung hiển thị nội dung lý thuyết.',
    'FormulaCard: thẻ hiển thị công thức hoặc ghi chú ngắn.',
    'Lesson2_TapHop: component chính của trang, quản lý ngôn ngữ và các trạng thái tương tác.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

# State management
p = doc.add_paragraph()
p.add_run('4. Quản lý trạng thái').bold = True
p = doc.add_paragraph('Trong component chính, code dùng React hooks để theo dõi các trạng thái sau:')
for item in [
    'lang: cho phép chuyển đổi giữa tiếng Việt và tiếng Anh.',
    'revealedAnswers: lưu tình trạng người dùng đã mở đáp án cho từng bài tập hay chưa.',
    'toggleAnswer: hàm thay đổi trạng thái mở/đóng đáp án.',
    't(vi, en): hàm tiện ích để lấy nội dung theo ngôn ngữ hiện tại.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

# Data structures
p = doc.add_paragraph()
p.add_run('5. Cấu trúc dữ liệu quan trọng').bold = True
p = doc.add_paragraph('Bài học dùng các mảng và object để quản lý nội dung một cách tách biệt khỏi giao diện:')
for item in [
    'learningObjectives: mục tiêu học tập.',
    'navItems: các mục điều hướng trong trang.',
    'videoSubtitles: phụ đề cho video, có thể bật theo từng từ khóa.',
    'mcQuestions: câu hỏi trắc nghiệm.',
    'tfCards: câu hỏi đúng/sai.',
    'fillQuestions: câu hỏi điền khuyết.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

# Render theory
p = doc.add_paragraph()
p.add_run('6. Chức năng renderTheory').bold = True
p = doc.add_paragraph('Hàm renderTheory là nơi xây dựng toàn bộ phần nội dung hiển thị trên trang. Nó trả về các section như:')
for item in [
    'Khởi động: đưa ra tình huống mở đầu.',
    'Video bài giảng: nhúng LessonVideoPlayer.',
    'Phần lý thuyết về tập hợp, tập con và hai tập bằng nhau.',
    'Phần thực hành: hiển thị bài tập và cho người dùng mở đáp án.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

# Premium engine
p = doc.add_paragraph()
p.add_run('7. Vai trò của PremiumLessonEngine').bold = True
p = doc.add_paragraph('File PremiumLessonEngine.js là lớp “bộ máy điều khiển” cho toàn bộ bài học. Nó không trực tiếp viết nội dung mà chịu trách nhiệm:')
for item in [
    'Tạo khung tổng thể của trang, gồm header, nội dung chính và các section.',
    'Quản lý logic làm bài: chọn đáp án, tính điểm, hiển thị kết quả.',
    'Hỗ trợ chế độ khó dễ của câu hỏi.',
    'Tích hợp gamification, animation và hiệu ứng reveal.',
    'Giao tiếp với auth và hệ thống lưu kết quả học tập.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

# How flow works
doc.add_paragraph('')
p = doc.add_paragraph()
p.add_run('8. Luồng hoạt động khi người dùng mở trang').bold = True
p = doc.add_paragraph('Khi trang được mở, hệ thống chạy theo trình tự sau:')
for item in [
    'Component Lesson2_TapHop khởi tạo trạng thái ngôn ngữ và các trạng thái đáp án.',
    'Nội dung bài học được truyền vào PremiumLessonEngine bằng props như lessonSlug, lessonTitle, mcQuestions, tfCards và fillQuestions.',
    'PremiumLessonEngine render giao diện chính và xử lý tương tác người dùng.',
    'Người dùng làm bài, hệ thống cập nhật điểm số và hiển thị kết quả.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

# Notes
doc.add_paragraph('')
p = doc.add_paragraph()
p.add_run('9. Những điểm cần chú ý khi mở rộng').bold = True
for item in [
    'Nên giữ tên id của các section thống nhất và không trùng nhau.',
    'Nếu thêm câu hỏi mới, cần đảm bảo dữ liệu có đúng định dạng với các trường q, options, answer, explain.',
    'Nếu muốn thêm ngôn ngữ mới, cần mở rộng cả nội dung và hàm t().',
    'Các phần tử có animation nên tránh lạm dụng quá nhiều để giữ trải nghiệm mượt.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(item)

# Footer
p = doc.add_paragraph()
p.add_run('Kết luận:').bold = True
p = doc.add_paragraph('Trang này là một ví dụ tốt về cách kết hợp React component, state management, dữ liệu nội dung và logic bài học thành một trải nghiệm web tương tác. Nếu cần, có thể tiếp tục mở rộng bằng cách thêm bài tập mới, tích hợp hệ thống chấm điểm tự động sâu hơn hoặc đổi giao diện sang dạng khóa học theo module.')

# Save
doc.save(out_path)
print(f'Created: {out_path}')
print(f'Exists: {out_path.exists()}')
