from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

out_dir = Path(r"D:\Downloads")
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / "Lesson2_TapHop_full_feature_explanation.docx"


doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

# Title
heading = doc.add_heading('Giải thích đầy đủ tất cả tính năng của trang bài học Tập Hợp', level=1)
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

p = doc.add_paragraph()
p.add_run('Tài liệu này giải thích toàn bộ chức năng của trang bài học được xây dựng từ hai file chính: ').bold = True
p.add_run('Lesson2_TapHop.js').italic = True
p.add_run(' và ')
p.add_run('PremiumLessonEngine.js').italic = True

p = doc.add_paragraph()
p.add_run('Ngày tạo: 2026-06-25').italic = True
p = doc.add_paragraph()

# Section 1
p = doc.add_paragraph()
p.add_run('1. Tổng quan về trang').bold = True
p = doc.add_paragraph()
p.add_run('Trang này là một bài học tương tác cho môn Toán 10, chuyên đề “Tập hợp”. Nó kết hợp nhiều tính năng: nội dung học, video bài giảng, bài tập trắc nghiệm, đúng/sai, điền khuyết, chuyển đổi ngôn ngữ, điều chỉnh độ khó, hiệu ứng chuyển động, gamification và lưu kết quả học tập.')

# Section 2
p = doc.add_paragraph()
p.add_run('2. Vai trò của file Lesson2_TapHop.js').bold = True
p = doc.add_paragraph('File này là file “nội dung” của bài học. Nó không xử lý toàn bộ logic mini-game, mà chịu trách nhiệm cung cấp dữ liệu và cấu trúc nội dung cho trang.')
for item in [
    'Định nghĩa các component nhỏ dùng để tạo khung UI: SectionHeader, TheoryBlock, FormulaCard.',
    'Tạo các biến dữ liệu cho bài học: lessonSlug, chapterTitle, lessonTitle, learningObjectives, navItems.',
    'Cung cấp dữ liệu video phụ đề bằng videoSubtitles.',
    'Tạo bộ câu hỏi: mcQuestions cho trắc nghiệm, tfCards cho đúng/sai, fillQuestions cho điền từ.',
    'Xây dựng toàn bộ nội dung lý thuyết bằng hàm renderTheory, chia thành các section: Khởi động, Video, lý thuyết về tập hợp, tập con, hai tập bằng nhau, thực hành.',
    'Quản lý trạng thái mở/đóng đáp án cho các bài tập thực hành bằng revealedAnswers và toggleAnswer.',
    'Cung cấp hàm t(vi, en) để đổi nội dung giữa tiếng Việt và tiếng Anh.',
]:
    para = doc.add_paragraph(style='List Bullet')
    para.add_run(item)

# Section 3
p = doc.add_paragraph()
p.add_run('3. Những tính năng cụ thể trong Lesson2_TapHop.js').bold = True
for item in [
    'Tạo header section đẹp và thống nhất cho từng phần bài học.',
    'Hiển thị định nghĩa, ví dụ và lưu ý về tập hợp, tập con và hai tập bằng nhau.',
    'Cho phép người dùng mở đáp án cho từng bài tập thực hành.',
    'Dùng các câu hỏi tương tác để kiểm tra kiến thức ngay trong trang.',
    'Cung cấp nội dung đa ngôn ngữ cho cả câu hỏi và giải thích.',
]:
    para = doc.add_paragraph(style='List Bullet')
    para.add_run(item)

# Section 4
p = doc.add_paragraph()
p.add_run('4. Vai trò của file PremiumLessonEngine.js').bold = True
p = doc.add_paragraph('File này là “bộ điều khiển trung tâm” của toàn bộ trang. Nó nhận nội dung từ Lesson2_TapHop.js, render giao diện và xử lý hầu hết logic tương tác.')
for item in [
    'Tạo thanh điều hướng sticky ở đầu trang để di chuyển nhanh giữa các section.',
    'Render tiêu đề trang, mục tiêu học tập và nội dung lý thuyết do Lesson2_TapHop truyền vào.',
    'Tạo khu vực mini-game ở cột phải, gồm 3 thể loại câu hỏi: trắc nghiệm, đúng/sai và điền từ.',
    'Quản lý trạng thái của các câu hỏi, điểm số và kết quả cuối cùng.',
    'Tích hợp animation, hiệu ứng reveal khi cuộn xuống, hiệu ứng nút bấm và particle burst.',
    'Tích hợp gamification, độ khó tự động/thủ công và lưu kết quả vào hệ thống người dùng.',
    'Kết nối với các component khác như DuoTranslate, MathToolsPanel, LessonVideoPlayer và useGamification.',
]:
    para = doc.add_paragraph(style='List Bullet')
    para.add_run(item)

# Section 5 - state
p = doc.add_paragraph()
p.add_run('5. Các trạng thái chính trong PremiumLessonEngine').bold = True
for item in [
    'lang: trạng thái ngôn ngữ hiện tại, dùng để chuyển đổi giữa tiếng Việt và tiếng Anh.',
    'gameMode: chọn chế độ câu hỏi hiện tại, gồm mc, tf và fill.',
    'activeSection: section đang được focus khi người dùng bấm vào mục điều hướng.',
    'difficultyMode: chế độ độ khó là auto hay manual.',
    'difficulty: mức độ khó hiện tại là NB, TH, VD hay VDC.',
    'mcIndex, mcSelected, mcScore, mcDone, mcHistory: theo dõi tiến trình câu hỏi trắc nghiệm.',
    'tfIndex, tfFlipped, tfScore, tfDone, tfHistory: theo dõi tiến trình câu hỏi đúng/sai.',
    'fillAnswers, fillChecked: lưu đáp án điền từ và trạng thái đã kiểm tra.',
]:
    para = doc.add_paragraph(style='List Bullet')
    para.add_run(item)

# Section 6 - features
p = doc.add_paragraph()
p.add_run('6. Giải thích từng tính năng').bold = True

p = doc.add_paragraph()
p.add_run('6.1 Thanh điều hướng sticky').bold = True
p = doc.add_paragraph('Thanh điều hướng ở đầu trang giúp người dùng di chuyển nhanh giữa các phần như Khởi động, Video, phần lý thuyết, Thực hành và Mini Game. Khi bấm vào một mục, hàm scrollTo sẽ gọi scrollIntoView để cuộn tới section tương ứng và cập nhật activeSection.')

p = doc.add_paragraph()
p.add_run('6.2 Chuyển đổi ngôn ngữ').bold = True
p = doc.add_paragraph('Nút VI/EN ở góc trên cùng cho phép đổi ngôn ngữ của toàn bộ bài học. Hàm t(vi, en) sẽ chọn nội dung phù hợp với lang hiện tại. Điều này giúp bài giảng có thể hiển thị bằng tiếng Việt hoặc tiếng Anh một cách đồng nhất.')

p = doc.add_paragraph()
p.add_run('6.3 Hiển thị mục tiêu học tập').bold = True
p = doc.add_paragraph('Phần mục tiêu học tập nằm ngay dưới tiêu đề chính. Nó dùng learningObjectives được truyền vào từ Lesson2_TapHop.js để hiện ra một danh sách ngắn gọn về những gì người học cần đạt được sau bài này.')

p = doc.add_paragraph()
p.add_run('6.4 Nội dung lý thuyết được inject vào trang').bold = True
p = doc.add_paragraph('RenderTheory là hàm do Lesson2_TapHop.js truyền vào. PremiumLessonEngine dùng props renderTheory để chèn toàn bộ nội dung bài học vào cột trái. Đây là cách tách riêng dữ liệu và giao diện, giúp code dễ mở rộng và tái sử dụng.')

p = doc.add_paragraph()
p.add_run('6.5 Mini Game ở cột phải').bold = True
p = doc.add_paragraph('Khu vực mini-game là phần tương tác chính của trang. Nó có 3 tab: Trắc nghiệm, Đúng/Sai và Điền từ. Người dùng chỉ cần chọn một chế độ để bắt đầu làm bài.')

p = doc.add_paragraph()
p.add_run('6.6 Chế độ độ khó tự động và thủ công').bold = True
p = doc.add_paragraph('Trang có hai chế độ độ khó. Khi ở chế độ auto, hệ thống gọi API getNextDifficulty để lấy đề xuất mức độ khó phù hợp với người dùng. Khi ở chế độ manual, người dùng có thể tự chọn NB, TH, VD hoặc VDC. Khi độ khó thay đổi, toàn bộ trạng thái câu hỏi được reset để bắt đầu lại từ đầu.')

p = doc.add_paragraph()
p.add_run('6.7 Chế độ trắc nghiệm').bold = True
p = doc.add_paragraph('Khi người dùng chọn một đáp án, hàm handleMcSelect sẽ ghi nhận lựa chọn. Nếu đúng thì điểm tăng. Sau đó nút tiếp theo sẽ hiện ra. Khi hết câu, hệ thống sẽ tính kết quả, gửi dữ liệu gamification và hiển thị ResultSummary.')

p = doc.add_paragraph()
p.add_run('6.8 Chế độ đúng/sai').bold = True
p = doc.add_paragraph('Mỗi thẻ câu hỏi hiển thị một mệnh đề. Người dùng chọn ĐÚNG hoặc SAI. Sau khi chọn, hệ thống sẽ hiện lời giải thích và cho phép chuyển sang câu tiếp theo. Điểm số được tính dựa trên số câu trả lời đúng.')

p = doc.add_paragraph()
p.add_run('6.9 Chế độ điền từ').bold = True
p = doc.add_paragraph('Ở chế độ này, người dùng nhập câu trả lời vào từng ô input. Khi nhấn kiểm tra, hệ thống sẽ so sánh đáp án với câu trả lời chuẩn, bỏ qua khoảng trắng và chữ hoa/thường. Nếu đúng, câu đó được tính là đúng.')

p = doc.add_paragraph()
p.add_run('6.10 Hệ thống gamification').bold = True
p = doc.add_paragraph('Mỗi khi người dùng hoàn thành một bộ câu hỏi, PremiumLessonEngine gọi triggerGamification để gửi dữ liệu về hệ thống như số câu đúng, thời gian làm bài, độ khó và độ chính xác. Nếu có HUD, giao diện gamification sẽ hiện lên để thông báo phần thưởng hoặc tiến độ.')

p = doc.add_paragraph()
p.add_run('6.11 Hiệu ứng animation').bold = True
p = doc.add_paragraph('Trang dùng Framer Motion để tạo animation cho nút, chuyển trạng thái câu hỏi và hiệu ứng hiện ra. Có cả ParticleBurst cho nút kiểm tra và MorphButton cho trải nghiệm tương tác với hiệu ứng khi bấm nút.')

p = doc.add_paragraph()
p.add_run('6.12 Kết quả và chơi lại').bold = True
p = doc.add_paragraph('Sau khi hoàn thành, hệ thống sẽ hiển thị tổng điểm và danh sách từng câu hỏi cùng đáp án đúng, đáp án của người dùng và trạng thái đúng/sai. Người dùng có thể bấm chơi lại để bắt đầu bộ câu hỏi mới từ đầu.')

p = doc.add_paragraph()
p.add_run('6.13 Lưu kết quả người dùng').bold = True
p = doc.add_paragraph('Khi người dùng hoàn thành một chế độ câu hỏi, hook saveGameResult sẽ được gọi để lưu tiến độ vào hệ thống auth. Mỗi lần có dữ liệu mới, giáo viên hoặc hệ thống có thể theo dõi tiến độ học tập của người dùng.')

p = doc.add_paragraph()
p.add_run('6.14 Các component bên ngoài tích hợp').bold = True
for item in [
    'DuoTranslate: hỗ trợ dịch thuật nhanh trong giao diện.',
    'MathToolsPanel: công cụ hỗ trợ toán học, hữu ích cho học sinh khi làm bài.',
    'LessonVideoPlayer: nhúng video bài giảng và phụ đề.',
    'useGamification: hook điều khiển trải nghiệm gamification.',
    'GamificationHUD: hiển thị thông báo phần thưởng và tiến độ.',
]:
    para = doc.add_paragraph(style='List Bullet')
    para.add_run(item)

# Section 7 - flow
p = doc.add_paragraph()
p.add_run('7. Luồng hoạt động khi người dùng mở trang').bold = True
for item in [
    'Lesson2_TapHop khởi tạo dữ liệu bài học, câu hỏi và nội dung lý thuyết.',
    'PremiumLessonEngine nhận các props và render giao diện chính.',
    'Người dùng chọn chế độ câu hỏi và độ khó.',
    'Hệ thống điều hướng và hiển thị câu hỏi tương ứng.',
    'Sau khi trả lời, điểm số và lời giải thích được cập nhật.',
    'Khi hoàn thành, kết quả được lưu và hiển thị.',
]:
    para = doc.add_paragraph(style='List Bullet')
    para.add_run(item)

# Section 8 - key takeaways
p = doc.add_paragraph()
p.add_run('8. Điểm cần ghi nhớ').bold = True
for item in [
    'Trang này là một ví dụ tốt về cách tách nội dung bài học và logic giao diện thành hai tầng rõ ràng.',
    'State management được dùng để điều khiển mọi trạng thái tương tác.',
    'Các tính năng như độ khó, gamification và lưu kết quả giúp tăng tính hấp dẫn và học tập hiệu quả.',
    'Nếu cần mở rộng, ta có thể thêm thêm loại câu hỏi mới, thêm ngôn ngữ mới hoặc đổi giao diện mà không làm ảnh hưởng đến logic cốt lõi.',
]:
    para = doc.add_paragraph(style='List Bullet')
    para.add_run(item)

# Conclusion
p = doc.add_paragraph()
p.add_run('Kết luận').bold = True
p = doc.add_paragraph('Trang này không chỉ là một trang hiển thị nội dung mà còn là một hệ thống học tập tương tác với nhiều tính năng: chuyển đổi ngôn ngữ, điều chỉnh độ khó, mini-game, animation, gamification, lưu kết quả và tích hợp nhiều công cụ học tập. Đây là một kiến trúc phù hợp để phát triển các bài học số một cách chuyên nghiệp.')

# Save

doc.save(out_path)
print(f'Created: {out_path}')
print(f'Exists: {out_path.exists()}')
